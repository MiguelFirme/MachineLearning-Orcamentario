"""Preenche a cópia do relatório acadêmico com os resultados do projeto."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


RAIZ = Path(__file__).resolve().parents[1]
ENTRADA = RAIZ / "Relatório_preenchido.docx"
SAIDA = RAIZ / "Relatório_preenchido.docx"
GRAFICOS = RAIZ / "graficos_treinamento"

AZUL = "1F4E79"
CINZA_CLARO = "EAF0F6"
BORDA = "D9D9D9"


def limpar_documento(doc):
    corpo = doc._element.body
    for elemento in list(corpo):
        if elemento.tag != qn("w:sectPr"):
            corpo.remove(elemento)


def definir_fonte(run, nome="Arial", tamanho=12, negrito=False, italico=False, cor="000000"):
    run.font.name = nome
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), nome)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), nome)
    run.font.size = Pt(tamanho)
    run.bold = negrito
    run.italic = italico
    run.font.color.rgb = RGBColor.from_string(cor)


def configurar_documento(doc):
    secao = doc.sections[0]
    secao.page_width = Cm(21)
    secao.page_height = Cm(29.7)
    secao.top_margin = Cm(3)
    secao.left_margin = Cm(3)
    secao.right_margin = Cm(2)
    secao.bottom_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    for nome, tamanho in [("Title", 18), ("Heading 1", 14), ("Heading 2", 12)]:
        estilo = doc.styles[nome]
        estilo.font.name = "Arial"
        estilo._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        estilo._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        estilo.font.size = Pt(tamanho)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        estilo.paragraph_format.space_before = Pt(12)
        estilo.paragraph_format.space_after = Pt(8)
        estilo.paragraph_format.keep_with_next = True

    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    rodape = secao.footer.paragraphs[0]
    rodape.clear()
    rodape.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    campo_inicio = OxmlElement("w:fldChar")
    campo_inicio.set(qn("w:fldCharType"), "begin")
    instrucao = OxmlElement("w:instrText")
    instrucao.set(qn("xml:space"), "preserve")
    instrucao.text = " PAGE "
    campo_fim = OxmlElement("w:fldChar")
    campo_fim.set(qn("w:fldCharType"), "end")
    run = rodape.add_run()
    run._r.extend([campo_inicio, instrucao, campo_fim])
    definir_fonte(run, tamanho=10)


def paragrafo(doc, texto="", negrito_inicial=None, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY,
              recuo=True, tamanho=12, italico=False):
    p = doc.add_paragraph()
    p.alignment = alinhamento
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.first_line_indent = Cm(1.25) if recuo else Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    if negrito_inicial and texto.startswith(negrito_inicial):
        r1 = p.add_run(negrito_inicial)
        definir_fonte(r1, tamanho=tamanho, negrito=True)
        r2 = p.add_run(texto[len(negrito_inicial):])
        definir_fonte(r2, tamanho=tamanho, italico=italico)
    else:
        r = p.add_run(texto)
        definir_fonte(r, tamanho=tamanho, italico=italico)
    return p


def titulo(doc, texto, nivel=1):
    p = doc.add_paragraph(style=f"Heading {nivel}")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(texto)
    definir_fonte(r, tamanho=14 if nivel == 1 else 12, negrito=True)
    return p


def lista(doc, itens, numerada=False):
    for indice, item in enumerate(itens, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        marcador = f"{indice}. " if numerada else "• "
        r = p.add_run(marcador + item)
        definir_fonte(r, tamanho=12)


def bordas_tabela(tabela):
    propriedades = tabela._tbl.tblPr
    bordas = propriedades.first_child_found_in("w:tblBorders")
    if bordas is None:
        bordas = OxmlElement("w:tblBorders")
        propriedades.append(bordas)
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:val"), "single")
        elemento.set(qn("w:sz"), "4")
        elemento.set(qn("w:color"), BORDA)
        bordas.append(elemento)


def preencher_celula(celula, texto, cabecalho=False, centralizar=False, alternada=False):
    celula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    tc_pr = celula._tc.get_or_add_tcPr()
    margem = OxmlElement("w:tcMar")
    for lado in ("top", "left", "bottom", "right"):
        item = OxmlElement(f"w:{lado}")
        item.set(qn("w:w"), "100")
        item.set(qn("w:type"), "dxa")
        margem.append(item)
    tc_pr.append(margem)
    if cabecalho or alternada:
        sombreado = OxmlElement("w:shd")
        sombreado.set(qn("w:fill"), AZUL if cabecalho else CINZA_CLARO)
        tc_pr.append(sombreado)
    p = celula.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centralizar or cabecalho else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(texto))
    definir_fonte(r, tamanho=10, negrito=cabecalho, cor="FFFFFF" if cabecalho else "000000")


def tabela(doc, cabecalhos, linhas, larguras=None, centralizadas=None):
    t = doc.add_table(rows=1, cols=len(cabecalhos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(layout)
    bordas_tabela(t)
    centralizadas = set(centralizadas or [])
    for indice, texto in enumerate(cabecalhos):
        preencher_celula(t.rows[0].cells[indice], texto, cabecalho=True)
    for numero, linha in enumerate(linhas):
        celulas = t.add_row().cells
        for indice, texto in enumerate(linha):
            preencher_celula(
                celulas[indice], texto, centralizar=indice in centralizadas, alternada=numero % 2 == 1
            )
    if larguras:
        for indice, largura in enumerate(larguras):
            t.columns[indice].width = Cm(largura)
            t._tbl.tblGrid.gridCol_lst[indice].set(qn("w:w"), str(int(Cm(largura).twips)))
        for linha in t.rows:
            for indice, largura in enumerate(larguras):
                celula = linha.cells[indice]
                celula.width = Cm(largura)
                tc_w = celula._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
                tc_w.set(qn("w:w"), str(int(Cm(largura).twips)))
                tc_w.set(qn("w:type"), "dxa")
    cabecalho_pr = t.rows[0]._tr.get_or_add_trPr()
    repetir = OxmlElement("w:tblHeader")
    repetir.set(qn("w:val"), "true")
    cabecalho_pr.append(repetir)
    for linha in t.rows:
        linha_pr = linha._tr.get_or_add_trPr()
        nao_dividir = OxmlElement("w:cantSplit")
        linha_pr.append(nao_dividir)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figura(doc, arquivo, legenda, largura=14.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.right_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(GRAFICOS / arquivo), width=Cm(largura))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    cap.paragraph_format.left_indent = Cm(0)
    cap.paragraph_format.right_indent = Cm(0)
    cap.paragraph_format.line_spacing = 1.0
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(legenda)
    definir_fonte(r, tamanho=10)


def quebra(doc):
    doc.add_page_break()


def construir_relatorio(doc):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(120)
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Previsão do Tempo de Corte Plasma por Geometria DXF com KNN e Regressão Ridge")
    definir_fonte(r, tamanho=18, negrito=True)
    paragrafo(
        doc,
        "Projeto final de aplicação de Machine Learning Clássico",
        alinhamento=WD_ALIGN_PARAGRAPH.CENTER,
        recuo=False,
        tamanho=13,
    )
    paragrafo(doc, "Setembro de 2026", alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=False)
    quebra(doc)

    titulo(doc, "Resumo executivo", 1)
    paragrafo(
        doc,
        "Este projeto prevê o tempo de corte plasma de peças metálicas a partir do material, da "
        "espessura e de atributos geométricos extraídos de arquivos DXF. Foram comparados o KNN "
        "Regressor, a regressão Ridge múltipla e a regressão Ridge polinomial de grau 2. A "
        "avaliação utilizou validação cruzada aninhada e agrupamento de entradas geométricas "
        "idênticas. O KNN apresentou o melhor resultado fora da amostra, com MAE de 17,43 s, "
        "RMSE de 27,51 s e R² de 0,722. Entre as regressões, a polinomial foi mantida como segundo "
        "modelo porque reduziu o MAE da múltipla em 2,54%."
    )
    paragrafo(
        doc,
        "Palavras chave: corte plasma; previsão de tempo; KNN; regressão Ridge; DXF; manufatura.",
        negrito_inicial="Palavras chave:", recuo=False,
    )

    titulo(doc, "1 Introdução problema e justificativa", 1)
    titulo(doc, "1.1 Contextualização do problema real", 2)
    paragrafo(
        doc,
        "O projeto está inserido no domínio de fabricação metalmecânica, especificamente no "
        "planejamento e na orçamentação do corte de chapas por plasma. Antes da fabricação, a "
        "empresa precisa estimar quanto tempo uma peça ocupará a máquina para calcular custos, "
        "organizar a carga produtiva e informar prazos. O tempo depende da espessura, do material, "
        "do comprimento total de corte e da complexidade geométrica, como furos, contornos internos "
        "e segmentos curvos."
    )
    paragrafo(
        doc,
        "A prática atual costuma combinar experiência dos responsáveis, tempos históricos e regras "
        "simples associadas a peso, perímetro ou espessura. Essa abordagem não considera de forma "
        "consistente todas as características da peça e pode produzir orçamentos diferentes para "
        "geometrias semelhantes. O erro afeta custo, prazo e alocação da máquina: uma previsão baixa "
        "subestima a capacidade necessária, enquanto uma previsão alta reduz a competitividade do orçamento."
    )
    paragrafo(
        doc,
        "Bakker (2023) aplicou regressão linear e estudo de tempos à previsão de carga de trabalho "
        "em uma máquina de plasma V310, mostrando a relação entre estimativas de tempo e planejamento "
        "operacional. Ma’ruf, Thoriq e Buwana (2024) compararam modelos de Machine Learning para "
        "estimativa antecipada do tempo de usinagem CNC e destacaram a influência do comprimento de "
        "corte, da espessura e de variáveis geométricas. Esses trabalhos apoiam o uso de dados "
        "históricos e características geométricas para estimar tempo de processamento."
    )

    titulo(doc, "1.2 Definição do problema de Machine Learning", 2)
    paragrafo(
        doc,
        "O problema é de regressão supervisionada porque a saída é um valor numérico contínuo, "
        "medido em segundos. A variável alvo é tempo_corte_plasma_s. Cada observação representa "
        "uma peça com tempo conhecido e características do material e da geometria."
    )
    paragrafo(
        doc,
        "O dataset utilizado é data/dataset_plasma_v10.csv, uma base interna consolidada a partir "
        "de registros do processo e informações geométricas extraídas dos desenhos DXF. Ele não "
        "provém de Kaggle ou UCI e acompanha o repositório entregue com o projeto. O arquivo bruto "
        "possui 1.919 linhas e 76 colunas. Após a aplicação dos critérios de qualidade e completude, "
        "1.141 peças foram usadas no treinamento. O conjunto é adequado porque associa o tempo real "
        "de plasma às entradas que estarão disponíveis quando uma nova peça for orçada."
    )

    titulo(doc, "1.3 Justificativa e relevância", 2)
    paragrafo(
        doc,
        "A solução fornece uma estimativa repetível antes da fabricação e pode apoiar orçamento, "
        "sequenciamento e análise de capacidade. O benefício esperado é reduzir decisões baseadas "
        "somente em julgamento individual e aproveitar o histórico de peças já processadas."
    )
    paragrafo(
        doc,
        "Machine Learning é pertinente porque o tempo resulta da interação de muitas características "
        "geométricas e de processo. Uma fórmula única baseada apenas no perímetro não representa "
        "adequadamente diferenças como quantidade de furos, comprimento curvo, espessura e material. "
        "Os modelos supervisionados aprendem essas relações a partir dos exemplos medidos e podem ser "
        "reavaliados quando novos registros forem incorporados."
    )

    titulo(doc, "1.4 Objetivos", 1)
    titulo(doc, "1.4.1 Objetivo geral", 2)
    paragrafo(
        doc,
        "Desenvolver e comparar modelos de Machine Learning Clássico para prever o tempo de corte "
        "plasma de peças metálicas a partir do material, da espessura e da geometria extraída de arquivos DXF."
    )
    titulo(doc, "1.4.2 Objetivos específicos", 2)
    lista(
        doc,
        [
            "Obter, inspecionar e estruturar o dataset do processo de corte plasma.",
            "Extrair dos arquivos DXF atributos de dimensão, perímetro, área, contornos e tipos de entidades.",
            "Diagnosticar valores ausentes, valores extremos, distribuição do alvo e repetição de geometrias.",
            "Construir pipelines de imputação, codificação categórica e padronização numérica.",
            "Treinar e otimizar KNN Regressor, regressão Ridge múltipla e regressão Ridge polinomial.",
            "Comparar os modelos por MAE, RMSE, R² e percentuais de erro em segundos.",
            "Disponibilizar um programa que receba um DXF, material e espessura e retorne as duas previsões finais.",
        ],
    )

    titulo(doc, "2 Fundamentação teórica", 1)
    titulo(doc, "2.1 Revisão bibliográfica do domínio", 2)
    paragrafo(
        doc,
        "O corte plasma é um processo térmico no qual um jato de gás ionizado em alta temperatura "
        "funde e remove material ao longo de uma trajetória. A duração de uma peça inclui o avanço "
        "ao longo do contorno, as mudanças de direção e os eventos associados a contornos internos. "
        "Neste projeto, tempo de corte significa o tempo registrado para a operação de plasma da peça; "
        "ele não deve ser confundido com horas de operador, tempo total do lote ou qualidade do kerf."
    )
    paragrafo(
        doc,
        "Salonitis e Vatousianos (2012) mostraram, por análise estatística e regressão, que corrente, "
        "velocidade, altura e pressão do gás influenciam respostas do corte plasma. O estudo trata "
        "qualidade do corte, e não tempo por peça, mas demonstra que modelos empíricos do processo "
        "dependem das condições operacionais e do material. O trabalho de Ma’ruf, Thoriq e Buwana "
        "(2024) é mais próximo do alvo deste projeto: a estimativa antecipada do tempo de usinagem "
        "melhora decisões de negociação, custo e planejamento, e variáveis geométricas contribuíram "
        "para a previsão."
    )
    paragrafo(
        doc,
        "Na amostra analítica deste projeto, o tempo varia de 4,53 s a 300 s, com média de 82,01 s "
        "e mediana de 60 s. Pelo critério do intervalo interquartil, 107 registros do alvo são valores "
        "extremos. Eles foram mantidos porque representam peças reais e porque remover casos longos "
        "reduziria a capacidade de avaliar o comportamento dos modelos nas peças mais exigentes."
    )
    figura(doc, "01_distribuicao_tempo_real.png", "Figura 1 Distribuição dos tempos reais das 1.141 peças válidas")
    paragrafo(
        doc,
        "Métodos não baseados em Machine Learning incluem tabelas de velocidade nominal, regras "
        "proporcionais ao perímetro, consulta manual de peças semelhantes e estimativa por especialista. "
        "Eles são úteis como referência, mas tendem a representar poucas variáveis de cada vez. O "
        "pipeline proposto reúne 28 atributos e aplica a mesma regra de previsão a todas as novas peças."
    )

    titulo(doc, "2.2 Referencial teórico em Machine Learning", 2)
    paragrafo(
        doc,
        "O KNN Regressor é um método baseado em instâncias. Para uma nova peça, ele identifica as "
        "observações mais próximas no espaço de atributos e combina seus tempos conhecidos. O modelo "
        "testou 3, 5, 7, 9, 11, 15, 21 e 31 vizinhos, com pesos uniformes ou inversamente relacionados "
        "à distância. Também foram comparadas as distâncias Manhattan, Euclidiana e Chebyshev. A "
        "distância Manhattan obteve o menor MAE interno médio, 18,92 s, contra 19,46 s da Euclidiana "
        "e 20,72 s da Chebyshev."
    )
    figura(doc, "07_comparacao_distancias_knn.png", "Figura 2 Comparação das métricas de distância avaliadas no KNN")
    paragrafo(
        doc,
        "A regressão Ridge estima o alvo por uma combinação das entradas e acrescenta uma penalização "
        "L2 aos coeficientes. Essa regularização é adequada quando existem atributos correlacionados, "
        "como perímetros, área, peso e dimensões. A versão múltipla usa os atributos transformados "
        "linearmente. A versão polinomial acrescenta termos de grau 2 e interações antes do ajuste "
        "Ridge, permitindo relações curvas, mas aumentando a complexidade. O hiperparâmetro alpha "
        "controla a força da regularização."
    )
    paragrafo(
        doc,
        "Não foi adotado modelo de ensemble. Árvores de decisão, Random Forest e métodos derivados "
        "foram excluídos do escopo desta etapa. A comparação concentra-se em KNN e regressões "
        "regularizadas, todos modelos clássicos sem árvores."
    )

    titulo(doc, "3 Metodologia e desenvolvimento", 1)
    titulo(doc, "3.1 Análise exploratória de dados", 2)
    paragrafo(
        doc,
        "O arquivo bruto possui 1.919 linhas e 76 colunas. O conjunto de treinamento mantém 1.141 "
        "linhas que apresentam alvo, todas as 28 entradas e status de fontes igual a ok. Na matriz "
        "analítica existem 27 entradas numéricas, uma entrada categórica e uma variável alvo numérica. "
        "Foram identificados 1.109 grupos de entradas distintas; 64 linhas pertencem a grupos com "
        "geometrias repetidas."
    )
    tabela(
        doc,
        ["Indicador", "Resultado"],
        [
            ["Linhas do arquivo bruto", "1.919"],
            ["Colunas do arquivo bruto", "76"],
            ["Peças válidas para treinamento", "1.141"],
            ["Entradas do modelo", "28"],
            ["Grupos geométricos distintos", "1.109"],
            ["Tempo mínimo mediano e máximo", "4,53 s  60 s  300 s"],
            ["Valores extremos no alvo pelo IQR", "107"],
        ],
        larguras=[10.0, 5.5],
        centralizadas=[1],
    )
    paragrafo(
        doc,
        "No arquivo bruto, as colunas formam uma base ampla que também contém informações não usadas "
        "pelo modelo; por isso existem valores ausentes fora do subconjunto analítico. Considerando "
        "alvo e entradas do modelo, foram encontradas 8.204 células ausentes antes do filtro. Após a "
        "seleção das 1.141 peças válidas, não restaram valores ausentes. O pipeline ainda contém "
        "imputação para tornar o processamento defensivo em futuras execuções."
    )
    paragrafo(
        doc,
        "Como o problema é de regressão, não existe desbalanceamento de classes. Há, porém, concentração "
        "do alvo em tempos padronizados e concentração de material: 829 das 1.141 peças são SAE 1020. "
        "Essa composição deve ser considerada ao aplicar o modelo a materiais pouco representados."
    )

    titulo(doc, "3.2 Pré processamento e engenharia de atributos", 2)
    paragrafo(
        doc,
        "Não foi usada uma única divisão fixa de treino e teste. A avaliação externa utiliza GroupKFold "
        "com cinco dobras, e a seleção interna utiliza GroupKFold com quatro dobras. Entradas idênticas "
        "recebem o mesmo identificador de grupo e permanecem na mesma dobra, evitando que uma peça "
        "duplicada seja usada para treinar e testar o modelo na mesma rodada."
    )
    paragrafo(
        doc,
        "As colunas numéricas passam por imputação pela mediana e StandardScaler. O material passa por "
        "imputação pela categoria mais frequente e One Hot Encoding, com tratamento de categorias "
        "desconhecidas. A padronização é essencial no KNN porque impede que atributos de grande escala, "
        "como área em milímetros quadrados, dominem o cálculo da distância. Na regressão polinomial, "
        "as entradas numéricas são padronizadas, expandidas para grau 2 e novamente padronizadas."
    )
    paragrafo(
        doc,
        "A engenharia de atributos começa no arquivo DXF. O extrator calcula comprimento, largura, "
        "área, perímetro total, perímetro externo e interno, área dos recortes, quantidade de contornos "
        "e entidades, comprimentos reto e curvo, dimensões dos recortes, razão de perímetro interno, "
        "ocupação do retângulo envolvente e compacidade. O peso líquido é estimado pela área, espessura "
        "e densidade do aço. Esses atributos representam tamanho e complexidade da trajetória de corte."
    )

    titulo(doc, "3.3 Implementação treinamento e otimização", 2)
    paragrafo(
        doc,
        "A implementação usa Python 3.14 e scikit-learn 1.9.1. Pandas realiza a leitura e a seleção dos "
        "dados; NumPy apoia os cálculos; os transformadores e modelos são encadeados em Pipeline e "
        "ColumnTransformer; e Joblib salva os modelos finais. O script principal é "
        "src/treinar_modelos_plasma.py."
    )
    paragrafo(
        doc,
        "O GridSearchCV escolhe a combinação com menor MAE dentro das dobras internas. No KNN, a grade "
        "contém 48 combinações de quantidade de vizinhos, pesos e distância. Na regressão Ridge múltipla, "
        "alpha assume 0,01, 0,1, 1, 10, 100 ou 1.000. Na polinomial, alpha assume 0,1, 1, 10, 100 ou "
        "1.000. Previsões negativas são limitadas a zero antes da avaliação, pois tempo negativo não "
        "possui interpretação física."
    )
    paragrafo(
        doc,
        "A validação é aninhada. A busca interna escolhe hiperparâmetros apenas com a parcela de treino "
        "da dobra externa. Em seguida, o modelo escolhido prevê a parcela externa, que não participou "
        "da seleção. Após as cinco rodadas, todas as previsões externas são reunidas para calcular as "
        "métricas. Por fim, uma nova busca usa todas as peças válidas e produz os modelos destinados ao uso."
    )

    titulo(doc, "3.4 Avaliação e análise de resultados", 2)
    paragrafo(
        doc,
        "Como o alvo é contínuo, métricas de classificação como acurácia, precisão, recall e F1 não se "
        "aplicam. O MAE foi adotado como critério principal porque expressa o erro típico diretamente em "
        "segundos. O RMSE dá peso maior aos erros grandes. O R² indica a proporção da variação do tempo "
        "explicada pelo modelo. Também foram calculados percentis do erro e a parcela de previsões com "
        "erro máximo de 10, 20 e 30 segundos."
    )
    tabela(
        doc,
        ["Modelo", "MAE s", "RMSE s", "R²", "Erro até 20 s"],
        [
            ["KNN", "17,43", "27,51", "0,722", "70,46%"],
            ["Ridge múltipla", "22,91", "32,55", "0,612", "58,28%"],
            ["Ridge polinomial", "22,33", "32,96", "0,602", "59,86%"],
        ],
        larguras=[5.2, 2.5, 2.5, 2.2, 3.1],
        centralizadas=[1, 2, 3, 4],
    )
    figura(doc, "02_comparacao_metricas_modelos.png", "Figura 3 Comparação quantitativa dos modelos em previsões fora da amostra")
    figura(doc, "03_real_vs_previsto_knn.png", "Figura 4 Tempos reais e previstos pelo KNN")
    figura(doc, "04_real_vs_previsto_regressao_multipla.png", "Figura 5 Tempos reais e previstos pela regressão Ridge múltipla")
    figura(doc, "05_real_vs_previsto_regressao_polinomial.png", "Figura 6 Tempos reais e previstos pela regressão Ridge polinomial")
    figura(doc, "08_faixas_de_erro_modelos.png", "Figura 7 Percentual de previsões dentro das faixas de erro")
    paragrafo(
        doc,
        "O KNN foi superior nos três indicadores principais e manteve 70,46% das previsões dentro de "
        "20 s do valor real. O modelo final usa sete vizinhos, distância Manhattan e pesos por distância. "
        "Esse resultado indica que peças próximas no espaço de atributos apresentam tempos mais "
        "semelhantes do que os estimados por uma única relação global."
    )
    paragrafo(
        doc,
        "Entre as regressões, a polinomial reduziu o MAE da múltipla em 2,54%, superando o limite de 2% "
        "definido antes da comparação. Por isso ela foi escolhida como segundo modelo, com alpha igual "
        "a 100. A regressão múltipla, porém, apresentou RMSE e R² ligeiramente melhores. Isso mostra "
        "uma compensação: a polinomial melhora o erro absoluto médio, mas comete alguns erros grandes, "
        "enquanto a múltipla é mais simples e ligeiramente mais estável nessas métricas."
    )
    paragrafo(
        doc,
        "A interpretabilidade do KNN é local: a previsão pode ser explicada pelas peças vizinhas, mas "
        "não por um coeficiente global. A Ridge múltipla possui coeficientes mais fáceis de inspecionar; "
        "a polinomial perde parte dessa clareza devido ao número de termos e interações. Para uso "
        "operacional, o sistema exibe KNN e regressão polinomial separadamente e não calcula média "
        "automática entre eles."
    )

    titulo(doc, "4 Conclusão produto final e cronograma", 1)
    titulo(doc, "4.1 Conclusão", 2)
    paragrafo(
        doc,
        "O projeto demonstrou que características disponíveis antes da fabricação podem ser usadas "
        "para estimar o tempo de corte plasma. O KNN alcançou o melhor desempenho fora da amostra e "
        "foi mantido como modelo principal. A regressão Ridge polinomial foi mantida como referência "
        "independente porque apresentou MAE menor do que a Ridge múltipla segundo a regra definida."
    )
    paragrafo(
        doc,
        "Os resultados não devem ser extrapolados sem validação para outras máquinas, condições de "
        "corte ou materiais pouco representados. A base é predominantemente composta por aço SAE 1020, "
        "e erros maiores permanecem nas peças de tempo elevado. As próximas melhorias recomendadas são "
        "registrar parâmetros efetivos da máquina, ampliar a quantidade de materiais e monitorar o erro "
        "após cada novo lote."
    )

    titulo(doc, "4.2 Produto final", 2)
    paragrafo(
        doc,
        "O produto final é um programa de linha de comando que recebe o caminho ou código de um DXF, "
        "o material e a espessura. O arquivo é processado pelo extrator geométrico, as 28 entradas são "
        "montadas na mesma ordem do treinamento e os dois modelos retornam previsões em segundos. Os "
        "resultados são exibidos separadamente para permitir comparação e análise técnica."
    )

    titulo(doc, "4.3 Entregáveis e reprodutibilidade", 2)
    tabela(
        doc,
        ["Entregável", "Localização ou especificação"],
        [
            ["Documento final", "Relatório_preenchido.docx"],
            ["Código fonte", "Pastas src, data, model e resultados do repositório entregue"],
            ["Modelos treinados", "model/modelo_knn_plasma.joblib e model/modelo_regressao_plasma.joblib"],
            ["Resultados", "resultados/relatorio_treinamento.json e previsoes_validacao_cruzada.csv"],
            ["Gráficos", "graficos_treinamento"],
            ["Apresentação", "Slides a serem preparados para a defesa oral"],
        ],
        larguras=[4.8, 10.7],
    )
    paragrafo(
        doc,
        "A execução não exige GPU. Um computador com processador de múltiplos núcleos, 8 GB de memória "
        "e espaço para o ambiente Python é suficiente para o conjunto atual. O ambiente validado usa "
        "Python 3.14.0, scikit-learn 1.9.1, pandas 3.0.6, NumPy 2.5.3, SciPy 1.18.1, Joblib 1.6.0, "
        "Matplotlib 3.11.2 e python-docx 1.2.0. As versões completas constam em requirements.txt."
    )
    paragrafo(doc, "Procedimento de instalação e treinamento", negrito_inicial="Procedimento de instalação e treinamento", recuo=False)
    lista(
        doc,
        [
            "Criar o ambiente com python -m venv .venv.",
            "Ativar o ambiente e executar python -m pip install -r requirements.txt.",
            "Executar python src/treinar_modelos_plasma.py para reproduzir a validação e os modelos.",
            "Executar python src/gerar_graficos_treinamento.py para atualizar as figuras.",
            "Executar python src/prever_tempo_plasma.py para prever uma nova peça.",
        ],
        numerada=True,
    )

    titulo(doc, "4.4 Cronograma", 2)
    tabela(
        doc,
        ["Etapa", "Situação"],
        [
            ["Consolidação e auditoria do dataset", "Concluída"],
            ["Extração e seleção dos atributos DXF", "Concluída"],
            ["Implementação e validação dos modelos", "Concluída"],
            ["Comparação de distâncias e geração dos gráficos", "Concluída"],
            ["Programa de previsão e teste funcional", "Concluída"],
            ["Revisão final do texto e preparação dos slides", "Próxima etapa"],
        ],
        larguras=[10.5, 5.0],
        centralizadas=[1],
    )

    titulo(doc, "Referências", 1)
    referencias = [
        "BAKKER, Mirjam. Workload Prediction for the V310: An application of linear regression and time study to make an estimation on required man hours. University of Twente, 2023. Disponível em: https://essay.utwente.nl/97072/. Acesso em: 26 set. 2026.",
        "MA’RUF, Anas; THORIQ, Dimas Ahmad; BUWANA, Kresna Surya. An Early Machining Time Estimation for Make-to-Order Manufacturing Using Machine Learning Approach. Procedia CIRP, v. 130, p. 106-111, 2024. DOI: https://doi.org/10.1016/j.procir.2024.10.063.",
        "SALONITIS, Konstantinos; VATOUSIANOS, S. Experimental Investigation of the Plasma Arc Cutting Process. Procedia CIRP, v. 3, p. 287-292, 2012. DOI: https://doi.org/10.1016/j.procir.2012.07.050.",
        "SCIKIT-LEARN DEVELOPERS. KNeighborsRegressor. Scikit-learn 1.9.1 documentation. Disponível em: https://scikit-learn.org/stable/modules/generated/sklearn.neighbors.KNeighborsRegressor.html. Acesso em: 26 set. 2026.",
        "SCIKIT-LEARN DEVELOPERS. Ridge. Scikit-learn 1.9.1 documentation. Disponível em: https://scikit-learn.org/stable/modules/generated/sklearn.linear_model.Ridge.html. Acesso em: 26 set. 2026.",
        "SCIKIT-LEARN DEVELOPERS. GridSearchCV e GroupKFold. Scikit-learn 1.9.1 documentation. Disponível em: https://scikit-learn.org/stable/modules/model_evaluation.html. Acesso em: 26 set. 2026.",
    ]
    for ref in referencias:
        p = paragrafo(
            doc,
            ref,
            recuo=False,
            tamanho=11,
            alinhamento=WD_ALIGN_PARAGRAPH.LEFT,
        )
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(10)


def main():
    doc = Document(ENTRADA)
    limpar_documento(doc)
    configurar_documento(doc)
    construir_relatorio(doc)
    doc.core_properties.title = "Previsão do Tempo de Corte Plasma por Geometria DXF"
    doc.core_properties.subject = "Projeto final de Machine Learning Clássico"
    doc.save(SAIDA)
    print(f"Relatório preenchido salvo em: {SAIDA}")


if __name__ == "__main__":
    main()
